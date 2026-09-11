#!/usr/bin/env python3
"""Validate the built deck: slide count, shape bounds, fonts, placeholder text."""
import sys
from pptx import Presentation
from pptx.util import Emu

deck = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
prs = Presentation(deck)
errors, warnings = [], []

assert len(prs.slides._sldIdLst) == 6, "deck must have exactly 6 slides"

BAD = ["lorem", "ipsum", "todo", "[insert", "your team name",
       "this slide layout", "@sih idea submission- template"]
FOOTER_Y = 6.86

for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        name = sh.name
        try:
            l, t = Emu(sh.left).inches, Emu(sh.top).inches
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
        except TypeError:
            continue
        # bounds — skip template chrome (footer bar, page number, footer text,
        # title-slide artwork)
        chrome = name.startswith(("Rectangle 8", "Rectangle 9",
                                  "Slide Number", "Footer", "Picture",
                                  "Freeform", "Rectangle 24"))
        if not chrome:
            if t + h > FOOTER_Y + 0.01:
                errors.append(f"slide {i}: {name} bottom {t+h:.2f} > {FOOTER_Y}")
            if l < -0.01 or l + w > 13.35:
                errors.append(f"slide {i}: {name} crosses horizontal edge")
        if sh.has_text_frame:
            txt = sh.text_frame.text.lower()
            for b in BAD:
                if b in txt and i != 1:
                    errors.append(f"slide {i}: {name} contains {b!r}")
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size and run.font.size.pt < 7 and run.text.strip():
                        errors.append(
                            f"slide {i}: {name} run below 7pt: {run.text[:30]!r}")

for e in errors:
    print("ERROR:", e)
for w in warnings:
    print("warn: ", w)
print("OK" if not errors else f"{len(errors)} errors")
sys.exit(1 if errors else 0)
